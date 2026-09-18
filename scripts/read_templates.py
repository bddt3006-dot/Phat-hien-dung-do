import sys, os
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document

BASE = 'C:/Users/trong/Downloads/Chuẩn bị tài liệu dự án theo QT 7 bước-20260916T062058Z-1-001/Chuẩn bị tài liệu dự án theo QT 7 bước'

FILES = [
    '0. Quy trình 7 bước và tài liệu cần thiết.docx',
    '01. Mẫu Báo cáo nghiên cứu khả thi.docx',
    '02. Mẫu Báo cáo phân tích thị trường.docx',
    '05. Mẫu Đặc tả yêu cầu của phần mềm.docx',
    '06. Mẫu_Kế hoạch kịch bản kiểm thử.docx',
    '07. Mẫu_Báo cáo kiểm thử.docx',
    '10a. Mẫu Báo cáo về phiên bản Alpha.docx',
    '10b. Mẫu Báo cáo về phiên bản Beta.docx',
    '10c. Mẫu Báo cáo về phiên bản Release.docx',
    '11. HDSD.docx',
    '12. Tài liệu BA.docx',
]

for fname in FILES:
    fp = os.path.join(BASE, fname)
    if not os.path.exists(fp):
        print(f'NOT FOUND: {fname}')
        continue
    doc = Document(fp)
    print(f'\n{"="*80}')
    print(f'FILE: {fname}')
    print(f'{"="*80}')
    count = 0
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            print(f'  [{para.style.name}] {t}')
            count += 1
        if count >= 60:
            print('  ... (truncated)')
            break
    for table in doc.tables:
        print(f'\n  [TABLE]')
        for row in table.rows[:5]:
            print('  | ' + ' | '.join([c.text.strip()[:30] for c in row.cells]))
        break
